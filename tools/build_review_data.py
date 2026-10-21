"""Build review/initial-data.js from raw exports, so the weekly review opens with data loaded.

The output holds customer and deal data, so it is git-ignored and must never
be committed to this (public) repo.

Usage:
    python3 -I tools/build_review_data.py OUT.js PIPELINE.xls [TRACKER.xlsx] [ACCOUNTS.xlsx]

PIPELINE  SFDC opportunity export with schedule lines (HTML-table .xls).
TRACKER   HK weekly tracker (.xlsx with POS_Raw, Dashboard, Pipeline sheets): POS rows,
          Q target and Match Keys.
ACCOUNTS  FY27 account list (.xlsx with Type, Account Name, Account ID USD, Region,
          FY26 RAD, FY27 Projection).
"""
import io
import json
import math
import os
import sys

import openpyxl
import pandas as pd

PIPE_COLS = ['First Ship Fiscal Quarter New', 'Opportunity Number', 'Opportunity Name', 'Sales Stage', 'Product Line',
             'Product Sub Line', 'Named Product Family', 'Product Name', 'Schedule Amount (converted)',
             'Primary Distributor', 'Primary Channel Partner', 'Channel Partner Siebel Account Number', 'Account Name',
             'Siebel Account Number', 'Commit Category', 'Application Vertical', 'Schedule Date']
POS_COLS = ['Invoice Fiscal Quarter', 'Invoice Fiscal Week', 'POS End User Name - Submitted', 'NVL End User',
            'POS Extended Cost (Net) (USD @ Planned Rate)', 'Submitting Account Name', 'POS Client Name - Submitted',
            'Product Line']
ACCT_COLS = ['Type', 'Account Name', 'Account ID USD', 'Region', 'FY26 RAD', 'FY27 Projection']


def clean(v):
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return ''
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def pipeline(path):
    raw = open(path, encoding='latin-1').read()
    d = pd.read_html(io.StringIO(raw))[0]
    cols = [c for c in PIPE_COLS if c in d.columns]
    return {'name': os.path.basename(path), 'cols': cols, 'rows': [[clean(v) for v in r] for r in d[cols].itertuples(index=False)]}


def tracker(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['POS_Raw']
    head_row = next(i for i, r in enumerate(ws.iter_rows(max_row=30, values_only=True), 1) if 'Invoice Fiscal Quarter' in r)
    hdr = [c.value for c in ws[head_row]]
    ix = {h: i for i, h in enumerate(hdr) if h}
    cols = [c for c in POS_COLS if c in ix]
    rows = [[clean(r[ix[c]]) for c in cols] for r in ws.iter_rows(min_row=head_row + 1, values_only=True) if r[ix['Invoice Fiscal Quarter']]]
    keys = {}
    if 'Pipeline' in wb.sheetnames:
        prow = list(wb['Pipeline'].iter_rows(values_only=True))
        h = next((i for i, r in enumerate(prow) if 'Match Key' in r), None)
        if h is not None:
            a, k = prow[h].index('Account'), prow[h].index('Match Key')
            keys = {r[a]: r[k] for r in prow[h + 1:] if r[a] and r[k]}
    target = wb['Dashboard']['C6'].value if 'Dashboard' in wb.sheetnames else None
    return {'name': os.path.basename(path) + ' (POS_Raw)', 'cols': cols, 'rows': rows}, keys, target


def accounts(path):
    a = pd.read_excel(path)
    rows = []
    for t, n, i, r, rad, pj in a[ACCT_COLS].itertuples(index=False):
        rows.append([t, n, int(i) if not pd.isna(i) else '', r, '' if pd.isna(rad) else rad, 0 if pd.isna(pj) else float(pj)])
    return {'name': os.path.basename(path), 'rows': rows}


def main(argv):
    if len(argv) < 3:
        sys.exit(__doc__)
    out = {'pipe': pipeline(argv[2])}
    for path in argv[3:]:
        names = openpyxl.load_workbook(path, read_only=True).sheetnames
        if 'POS_Raw' in names:
            out['pos'], out['keys'], out['target'] = tracker(path)
        else:
            out['acct'] = accounts(path)
    js = 'window.HKPR_SAMPLE = ' + json.dumps(out, ensure_ascii=False, separators=(',', ':')) + ';\n'
    open(argv[1], 'w', encoding='utf-8').write(js)
    print(f"wrote {argv[1]}: {len(out['pipe']['rows'])} pipeline rows")


if __name__ == '__main__':
    main(sys.argv)
